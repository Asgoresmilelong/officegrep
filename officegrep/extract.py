"""从各类文档里抽出可检索的文本块。

每个 block 都带一个人类可读的定位符（段落号 / 页码 / 幻灯片号 / 单元格），
因为 agent 拿到搜索结果后第一个问题永远是「在哪儿」。
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

SUPPORTED = {".docx", ".pptx", ".xlsx", ".pdf", ".txt", ".md", ".markdown"}
MAX_BLOCK = 2000


@dataclass
class Block:
    loc: str
    text: str


def _clean(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract(path: str | Path) -> list[Block]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx(path)
    if suffix == ".pptx":
        return _pptx(path)
    if suffix == ".xlsx":
        return _xlsx(path)
    if suffix == ".pdf":
        return _pdf(path)
    if suffix in {".txt", ".md", ".markdown"}:
        return _plain(path)
    raise ValueError(f"不支持的格式：{suffix}")


# ------------------------------------------------------------------ docx


def _qn(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


def _docx_paragraphs(container, in_table: bool = False):
    """按文档顺序产出 (w:p, in_table)，下钻表格与内容控件。"""
    for el in container:
        if el.tag == _qn(W, "p"):
            yield el, in_table
        elif el.tag == _qn(W, "tbl"):
            for row in el.findall(_qn(W, "tr")):
                for cell in row.findall(_qn(W, "tc")):
                    yield from _docx_paragraphs(cell, True)
        elif el.tag in (_qn(W, "sdt"), _qn(W, "sdtContent")):
            content = el.find(_qn(W, "sdtContent")) if el.tag == _qn(W, "sdt") else el
            if content is not None:
                yield from _docx_paragraphs(content, in_table)


def _paragraph_text(p) -> str:
    parts = []
    for node in p.iter():
        if node.tag == _qn(W, "t") and node.text:
            parts.append(node.text)
        elif node.tag == _qn(W, "tab"):
            parts.append("\t")
        elif node.tag in (_qn(W, "br"), _qn(W, "cr")):
            parts.append("\n")
    return "".join(parts)


def _docx(path: Path) -> list[Block]:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml")
    body = etree.fromstring(xml).find(_qn(W, "body"))
    if body is None:
        return []

    blocks: list[Block] = []
    for idx, (p, in_table) in enumerate(_docx_paragraphs(body), start=1):
        text = _clean(_paragraph_text(p))
        if len(text) < 2:
            continue
        loc = f"{'table ' if in_table else ''}p {idx}"
        blocks.append(Block(loc=loc, text=text[:MAX_BLOCK]))
    return blocks


# ------------------------------------------------------------------ pptx


def _pptx(path: Path) -> list[Block]:
    with zipfile.ZipFile(path) as zf:
        names = [n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        names.sort(key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[-1]).group(1)))
        blocks: list[Block] = []
        for name in names:
            number = int(re.search(r"(\d+)", name.rsplit("/", 1)[-1]).group(1))
            root = etree.fromstring(zf.read(name))
            texts = [t.text for t in root.iter(_qn(A, "t")) if t.text]
            text = _clean(" ".join(texts))
            if len(text) >= 2:
                blocks.append(Block(loc=f"slide {number}", text=text[:MAX_BLOCK]))
    return blocks


# ------------------------------------------------------------------ xlsx


def _xlsx(path: Path) -> list[Block]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        for ws in wb.worksheets:
            for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if not cells:
                    continue
                text = _clean(" | ".join(cells))
                blocks.append(Block(loc=f"{ws.title}!{row_no}", text=text[:MAX_BLOCK]))
    finally:
        wb.close()
    return blocks


# ------------------------------------------------------------------ pdf


def _pdf(path: Path) -> list[Block]:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - 可选依赖
        raise RuntimeError("读取 PDF 需要安装 pypdf：pip install officegrep[pdf]")

    reader = PdfReader(str(path))
    blocks: list[Block] = []
    for i, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if len(text) >= 2:
            blocks.append(Block(loc=f"page {i}", text=text[:MAX_BLOCK]))
    return blocks


# ------------------------------------------------------------------ txt/md


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _plain(path: Path) -> list[Block]:
    text = _read_text(path).replace("\r\n", "\n")
    blocks: list[Block] = []
    start = 0
    buffer: list[str] = []
    length = 0

    def flush(line_no: int) -> None:
        nonlocal buffer, length, start
        if buffer:
            joined = _clean("\n".join(buffer))
            if len(joined) >= 2:
                blocks.append(Block(loc=f"line {start}", text=joined[:MAX_BLOCK]))
        buffer, length, start = [], 0, line_no

    for line_no, line in enumerate(text.split("\n"), start=1):
        if not buffer:
            start = line_no
        buffer.append(line)
        length += len(line)
        if length >= 800 or not line.strip():
            flush(line_no + 1)
    flush(len(text.split("\n")) + 1)
    return blocks

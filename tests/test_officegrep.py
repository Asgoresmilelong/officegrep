"""端到端测试：夹具用 python-docx / python-pptx / openpyxl 现场生成。"""

from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document as Docx
from openpyxl import Workbook
from pptx import Presentation

from officegrep import Index, extract
from officegrep.mcp import Server


class Base(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)

    def make_docx(self, name: str = "报告.docx") -> Path:
        doc = Docx()
        doc.add_paragraph("第一章 结构设计")
        doc.add_paragraph("复合材料电线杆在岛礁环境下采用变截面薄壁结构。")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格中的弯矩设计值"
        path = self.root / name
        doc.save(path)
        return path

    def make_pptx(self, name: str = "汇报.pptx") -> Path:
        prs = Presentation()
        for title, body in (("总体方案", "杆身分段成型"), ("成本分析", "出厂价 2324 元")):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        path = self.root / name
        prs.save(path)
        return path

    def make_xlsx(self, name: str = "参数表.xlsx") -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "材料"
        ws.append(["材料", "弹性模量", "强度"])
        ws.append(["GFRP", 28, 320])
        path = self.root / name
        wb.save(path)
        return path

    def make_txt(self, name: str = "笔记.md") -> Path:
        path = self.root / name
        path.write_text("第一段：弯矩分布\n\n第二段：挠度校核\n", encoding="utf-8")
        return path


class TestExtract(Base):
    def test_docx_paragraphs_and_table(self):
        blocks = extract(self.make_docx())
        text = " ".join(b.text for b in blocks)
        self.assertIn("复合材料电线杆", text)
        self.assertIn("表格中的弯矩设计值", text)
        self.assertTrue(any(b.loc.startswith("table p ") for b in blocks), [b.loc for b in blocks])

    def test_pptx_slides(self):
        blocks = extract(self.make_pptx())
        self.assertEqual([b.loc for b in blocks], ["slide 1", "slide 2"])
        self.assertIn("出厂价", blocks[1].text)

    def test_xlsx_rows(self):
        blocks = extract(self.make_xlsx())
        self.assertEqual(blocks[0].loc, "材料!1")
        self.assertIn("GFRP", " ".join(b.text for b in blocks))

    def test_plain_text(self):
        blocks = extract(self.make_txt())
        self.assertTrue(blocks[0].loc.startswith("line "))
        self.assertIn("弯矩", blocks[0].text)

    def test_unsupported_suffix(self):
        path = self.root / "x.zip"
        path.write_bytes(b"PK")
        with self.assertRaises(ValueError):
            extract(path)


class TestIndex(Base):
    def test_search_chinese_long_and_short(self):
        self.make_docx()
        with Index(self.root / "i.db") as idx:
            stats = idx.index_paths([self.root])
            self.assertEqual(stats["indexed"], 1)
            long_hits = idx.search("变截面薄壁结构")
            self.assertTrue(long_hits)
            self.assertIn("报告.docx", long_hits[0].path)
            short_hits = idx.search("弯矩")
            self.assertTrue(short_hits)
            self.assertIn("弯矩", short_hits[0].snippet)

    def test_incremental_skips_unchanged(self):
        self.make_docx()
        with Index(self.root / "i.db") as idx:
            first = idx.index_paths([self.root])
            second = idx.index_paths([self.root])
            self.assertEqual(first["indexed"], 1)
            self.assertEqual(second["indexed"], 0)
            self.assertEqual(second["skipped"], 1)

    def test_modified_file_is_reindexed(self):
        path = self.make_docx()
        with Index(self.root / "i.db") as idx:
            idx.index_paths([self.root])
            doc = Docx()
            doc.add_paragraph("新增内容：风荷载校核")
            doc.save(path)
            stats = idx.index_paths([self.root])
            self.assertEqual(stats["indexed"], 1)
            self.assertTrue(idx.search("风荷载"))

    def test_prune_removes_missing(self):
        path = self.make_docx()
        with Index(self.root / "i.db") as idx:
            idx.index_paths([self.root])
            path.unlink()
            self.assertEqual(idx.prune(self.root), 1)
            self.assertEqual(idx.search("弯矩"), [])

    def test_under_filter(self):
        self.make_docx()
        sub = self.root / "sub"
        sub.mkdir()
        other = self.root / "other.md"
        other.write_text("另一份文档里的弯矩", encoding="utf-8")
        with Index(self.root / "i.db") as idx:
            idx.index_paths([self.root])
            hits = idx.search("弯矩", under=sub)
            self.assertEqual(hits, [])


class TestMcp(Base):
    def setUp(self) -> None:
        super().setUp()
        self.make_docx()
        self.index = Index(self.root / "i.db")
        self.addCleanup(self.index.close)
        self.index.index_paths([self.root])
        self.server = Server(self.index)

    def call(self, method: str, params=None, msg_id=1):
        return self.server.handle({"jsonrpc": "2.0", "id": msg_id, "method": method, "params": params or {}})

    def test_initialize(self):
        resp = self.call("initialize", {"protocolVersion": "2025-06-18", "capabilities": {}})
        self.assertEqual(resp["result"]["serverInfo"]["name"], "officegrep")
        self.assertIn("tools", resp["result"]["capabilities"])

    def test_notification_has_no_response(self):
        self.assertIsNone(self.server.handle({"jsonrpc": "2.0", "method": "notifications/initialized"}))

    def test_tools_list(self):
        names = [t["name"] for t in self.call("tools/list")["result"]["tools"]]
        self.assertEqual(names, ["search_docs", "list_docs", "index_folder"])

    def test_tools_call_search(self):
        resp = self.call("tools/call", {"name": "search_docs", "arguments": {"query": "弯矩"}})
        text = resp["result"]["content"][0]["text"]
        self.assertIn("报告.docx", text)
        self.assertFalse(resp["result"]["isError"])

    def test_unknown_tool(self):
        resp = self.call("tools/call", {"name": "nope", "arguments": {}})
        self.assertEqual(resp["error"]["code"], -32602)

    def test_unknown_method(self):
        resp = self.call("resources/list")
        self.assertEqual(resp["error"]["code"], -32601)

    def test_stdio_loop(self):
        stdin = io.StringIO(
            json.dumps({"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {}}) + "\n"
            + json.dumps({"jsonrpc": "2.0", "id": 2, "method": "tools/list"}) + "\n"
        )
        stdout = io.StringIO()
        self.server.serve(stdin=stdin, stdout=stdout)
        lines = [json.loads(line) for line in stdout.getvalue().splitlines()]
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[1]["id"], 2)


class TestCli(Base):
    def test_cli_index_search(self):
        self.make_docx()
        db = self.root / "cli.db"
        from officegrep.cli import main

        self.assertEqual(main(["--db", str(db), "index", str(self.root)]), 0)
        self.assertEqual(main(["--db", str(db), "search", "弯矩"]), 0)
        self.assertEqual(main(["--db", str(db), "stats"]), 0)


if __name__ == "__main__":
    unittest.main()
